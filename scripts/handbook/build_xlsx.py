"""Build docs/handbook/sbmig-reference.xlsx from data.json (system python3 + openpyxl)."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="2C5F2D")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="2C5F2D")
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="555555")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, note, columns, rows, widths):
    """columns: [(header, key)], rows: list of dicts. Row 1 title, row 2 note, row 4 header, data from row 5."""
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE_FONT
    ws["A2"] = note; ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(columns)))
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    hr = 4
    for i, (h, _) in enumerate(columns, 1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, hr + 1):
        for i, (_, k) in enumerate(columns, 1):
            v = row.get(k, "") if isinstance(row, dict) else row[i - 1]
            c = ws.cell(row=r, column=i, value=v if v != "" else None)
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(columns))}{hr + len(rows)}"
    return ws, hr, len(rows)

specs = []   # (sheet name, description, header row, row count) for the index formulas

# ---- 설정키
ws, hr, n = sheet("설정키", "설정 키 레퍼런스 (sbmig.core.config.Config 전체)",
    "출처: pydantic 모델 Config의 모든 필드(코드에서 추출) + config/default.yml의 주석. '기본값(코드)'은 모델 기본값, 'default.yml 값'은 기본 설정 파일에 적힌 값입니다. "
    "회사 설정은 config/artifacts.yml(예시: config/artifacts.example.yml)에서 덮어씁니다. 상세 설명: docs/handbook/08-config-reference.md",
    [("섹션", "section"), ("키", "key"), ("타입", "type"), ("기본값(코드)", "default"), ("default.yml 값", "yml"), ("설명(default.yml 주석 / 필드 설명)", "comment")],
    [dict(r, section=r["key"].split(".")[0]) for r in data["config"]], [14, 44, 34, 30, 22, 90])
specs.append(("설정키", "모든 설정 키: 타입·기본값·설명", hr, n))

# ---- 핸들러
ws, hr, n = sheet("핸들러", "결정적 핸들러 카탈로그 (LLM 없이 고치는 수정기)",
    "출처: sbmig.triage.handlers.HandlerRegistry.default() 등록 순서 = 반응형 우선순위(먼저 맞는 핸들러가 이김). '선제 적용'은 첫 빌드 전에 프로필을 보고 미리 돌리는지(항상 / 프로필 기반 / 반응형만). "
    "'트리거'는 빌드 에러 문구에서 찾는 정규식입니다. 상세: docs/handbook/04-triage-and-loop.md 5장, docs/ERROR-HANDLING.md",
    [("순서", "n"), ("핸들러 id", "id"), ("하는 일", "description"), ("대상 에러 종류", "kinds"), ("선제 적용", "preemptive"), ("선제 단계", "phase"), ("트리거(정규식)", "triggers"), ("클래스", "class")],
    [dict(h, n=i) for i, h in enumerate(data["handlers"], 1)], [6, 34, 60, 34, 12, 10, 70, 50])
specs.append(("핸들러", "결정적 핸들러 23종: 트리거·선제 조건·하는 일", hr, n))

# ---- 에러카탈로그
cols = ["section"] + [k for k in ("ID", "Symptom", "Symptom (detection)", "Policy", "Report says", "Fixture", "Advisory", "Library", "Status", "Replacement", "Notes") if any(k in r for r in data["errors"])]
extra = sorted({k for r in data["errors"] for k in r} - set(cols))
cols += extra
ws, hr, n = sheet("에러카탈로그", "에러·엣지 케이스 카탈로그 (docs/ERROR-HANDLING.md 원문 표)",
    "출처: docs/ERROR-HANDLING.md의 표를 섹션별로 그대로 옮긴 것(영문). ID 접두: X = 외부/인프라 blocked, E = 환경·레거시 quirk(핸들러), L = 루프 안전장치, S = 구조 엣지, §6 = 라이브러리 어드바이저리. "
    "한국어 증상→의미→조치 설명: docs/handbook/09-operations.md 6장",
    [(c if c != "section" else "섹션", c) for c in cols], data["errors"], [34] + [14 if c == "ID" else 60 for c in cols[1:]])
specs.append(("에러카탈로그", "X/E/L/S 카탈로그 원문 63행", hr, n))

# ---- CLI
ws, hr, n = sheet("CLI", "CLI 명령과 옵션 (typer 앱에서 추출)",
    "출처: sbmig.cli.app (typer) → click 명령 트리. 한 행 = 한 옵션/인자. 사용 예와 설명: docs/handbook/07-api-reference.md 1장",
    [("명령", "command"), ("명령 설명", "help"), ("파라미터", "param"), ("옵션 표기", "opts"), ("타입", "type"), ("필수", "required"), ("기본값", "default"), ("파라미터 설명", "param_help")],
    data["cli"], [28, 50, 20, 30, 14, 6, 24, 70])
specs.append(("CLI", "sbmig 명령·옵션 전체", hr, n))

# ---- REST_API
ws, hr, n = sheet("REST_API", "웹 UI REST API (FastAPI 라우트에서 추출)",
    "출처: sbmig.web.app.create_app()의 라우트. 서버 기동: `sbmig serve` 또는 `uvicorn sbmig.web.app:app`. 이벤트 스트림(SSE)의 이벤트 종류와 페이로드: docs/handbook/07-api-reference.md 2장",
    [("메서드", "method"), ("경로", "path"), ("핸들러 이름", "name"), ("설명", "summary"), ("상태 코드", "status")],
    data["rest"], [10, 36, 26, 80, 10])
specs.append(("REST_API", "웹 REST 라우트 14개", hr, n))

# ---- 픽스처
ws, hr, n = sheet("픽스처", "e2e 픽스처 프로젝트 (fixtures/projects/*/expected.json)",
    "출처: 각 픽스처의 README.md 제목과 expected.json. e2e: scripts/e2e.sh (실제 Maven, JDK 21 필요). '최대 LLM 호출'이 0이면 결정적으로만 끝나야 하는 시나리오입니다. 상세: docs/handbook/05-directories-and-storage.md",
    [("이름", "name"), ("시나리오", "title"), ("케이스", "case"), ("Java", "java"), ("Boot", "boot"), ("Spring", "spring"), ("패키징", "packaging"), ("멀티모듈", "multi"),
     ("감지 라이브러리", "libraries"), ("기대 성공", "success"), ("기대 exit", "exit"), ("최대 LLM 호출", "max_llm"), ("기대 blocked", "blocked"), ("카탈로그(기대 어드바이저리 등)", "catalog")],
    data["fixtures"], [26, 70, 14, 8, 8, 8, 8, 8, 34, 8, 8, 8, 30, 40])
specs.append(("픽스처", "24개 픽스처와 기대 결과", hr, n))

# ---- 저널이벤트
ws, hr, n = sheet("저널이벤트", "journal.jsonl 이벤트 종류 (코드에서 추출)",
    "출처: sbmig/ 아래 `\"event\": \"…\"` 리터럴. 실행 디렉터리 out/<run id>/journal.jsonl에 한 줄 JSON으로 기록되며 웹 UI가 그대로 스트리밍합니다. 의미와 순서: docs/handbook/01-pipeline.md",
    [("이벤트", "event"), ("기록하는 모듈", "files")], data["events"], [34, 90])
specs.append(("저널이벤트", "journal.jsonl 이벤트 이름과 기록 모듈", hr, n))

# ---- 리포트플래그
ws, hr, n = sheet("리포트플래그", "report.json flags 값 (코드에서 추출)",
    "출처: `_flag(\"…\")` / `flags.append(\"…\")` / `.flag(\"…\")` 리터럴(f-string 접두는 `{…}` 자리에 실제 값이 들어감). 플래그는 exit code를 바꾸지 않는 '주의 표시'입니다. 뜻: docs/handbook/09-operations.md 3장",
    [("플래그", "flag"), ("설정하는 모듈", "files")], data["flags"], [44, 90])
specs.append(("리포트플래그", "report.flags 값 목록", hr, n))

# ---- 레시피카탈로그
ws, hr, n = sheet("레시피카탈로그", "OpenRewrite 레시피 카탈로그 (knowledge/recipes/catalog.yml)",
    "출처: knowledge/recipes/catalog.yml — 레시피 id의 단일 출처(불변식 5). `sbmig recipes check`가 config/default.yml의 레시피 세트, sbmig/triage/recipe_rules.yml과의 드리프트를 검사합니다. "
    "'advisor 허용'은 LLM 플랜 어드바이저가 추가할 수 있는지. 상세: docs/RECIPES.md",
    [("레시피 id", "id"), ("출처", "origin"), ("아티팩트", "artifact"), ("기본 단계", "phase"), ("advisor 허용", "advisor"), ("게이트", "gate"), ("세트", "sets"), ("규칙", "rules"), ("검증 시점", "verified"), ("설명", "description")],
    data["recipes"], [58, 10, 44, 10, 8, 18, 22, 22, 10, 80])
specs.append(("레시피카탈로그", "레시피 33개", hr, n))

# ---- 어드바이저리
ws, hr, n = sheet("어드바이저리", "라이브러리 어드바이저리 (knowledge/libraries/advisories.yml)",
    "출처: knowledge/libraries/advisories.yml. 정책은 guide(안내만) | recipe(레시피로 전환)뿐이며 block은 없습니다(회사 결정). 리포트 Handover에 모입니다. 상세: docs/ADVISORIES.md, docs/handbook/02-workflows.md 8장",
    [("id", "id"), ("라이브러리 규칙", "library"), ("아티팩트", "artifacts"), ("상태", "status"), ("정책", "policy"), ("제목", "title"), ("대체", "replacement"), ("대체 버전", "replacement_version"), ("레시피", "recipes"), ("적용 케이스", "applies_to"), ("요약", "summary")],
    data["advisories"], [16, 14, 34, 12, 8, 44, 40, 14, 44, 16, 90])
specs.append(("어드바이저리", "어드바이저리 14건", hr, n))

# ---- 지식노트
ws, hr, n = sheet("지식노트", "LLM 프롬프트용 지식 노트 (knowledge/notes/index.yml)",
    "출처: knowledge/notes/index.yml (각 노트 front-matter와 동기화, 테스트로 검사). 에러 문구·케이스·단계가 키워드와 맞는 노트가 수정 LLM의 프롬프트에 들어갑니다(sbmig/fix/context.py). 'always'는 항상 포함.",
    [("파일", "file"), ("제목", "title"), ("키워드(정규식)", "keywords"), ("케이스", "cases"), ("단계", "phases"), ("항상 포함", "always")],
    data["notes"], [26, 50, 90, 30, 20, 8])
specs.append(("지식노트", "노트 14편과 선택 키워드", hr, n))

# ---- 열거형
ws, hr, n = sheet("열거형", "분류 체계 (sbmig/core/models.py 열거형)",
    "출처: CaseType(프로젝트 유형), Phase(단계), StepKind(스텝 종류), ErrorKind(빌드 에러 종류), DecisionKind(트리아지 결정), PropertyStatus(외부화 설정 판정). 용어 설명: docs/handbook/README.md 5장",
    [("열거형", "enum"), ("값", "value"), ("비고", "note")], data["enums"], [20, 34, 14])
specs.append(("열거형", "케이스·단계·에러종류·결정종류 값", hr, n))

# ---- 디렉터리 (hand-written)
dirs = [
 ("config/", "설정 yml, Maven settings 템플릿", "사람(운영자)", "load_config (모든 명령)", "버전관리. 키·토큰은 절대 넣지 않음(.env)", "default.yml → --config 파일 순서 → .env → 쉘 환경 → CLI"),
 ("config/default.yml", "모든 키의 기본값(코드와 함께 배포)", "개발자", "load_config", "수정 대신 artifacts.yml에서 덮어쓰기", ""),
 ("config/artifacts.example.yml", "회사 아티팩트 자리표시자(xxx-parent 등), parent_phase: spring", "운영자가 복사해 artifacts.yml 작성", "load_config", "artifacts.yml은 회사 값이므로 사내에서만 관리", ""),
 ("config/artifacts.local-e2e.yml", "공개 대체 아티팩트 + 9c 이전 레이아웃 고정(옛 픽스처용)", "개발자", "scripts/e2e.sh, 테스트", "회사에서는 쓰지 않음", ""),
 ("config/settings.nexus.example.xml", "Nexus catch-all 미러 settings.xml 템플릿", "운영자가 복사", "build.settings_file로 모든 Maven 호출", "URL만 바꾸면 됨", "RUNBOOK §0.1"),
 ("knowledge/recipes/catalog.yml", "레시피 id 단일 출처", "개발자(라운드마다)", "플래너 어드바이저, sbmig recipes check", "버전관리", "불변식 5"),
 ("knowledge/recipes/base/*.yml", "선언적 OpenRewrite 레시피(사내 전용: 버전 핀, Hystrix→Resilience4j 등)", "개발자", "rewrite 실행 시 compose", "버전관리", ""),
 ("knowledge/notes/*.md + index.yml", "LLM 프롬프트용 마이그레이션 노트(키워드 선택)", "개발자", "sbmig/fix/context.py", "버전관리, index.yml과 front-matter 동기화 테스트", ""),
 ("knowledge/libraries/advisories.yml", "EOL/비관리 라이브러리 안내(guide|recipe)", "개발자", "sbmig/advisory", "회사 추가: advisories.extra_files", ""),
 ("knowledge/config/externalized.yml", "외부화 설정 인벤토리 규칙(JNDI → spring.datasource.* 등)", "개발자", "sbmig/analyze/properties.py", "버전관리", ""),
 ("kb/ledger.jsonl", "모든 시도 기록(append-only): 시그니처 → 수정 라벨 → 결과", "실행(fix loop, harvest)", "플릿 요약, 사람 검토, 레시피 승격", "학습 자산 — 백업·공유. 실행 간 공유 디렉터리로 운영(kb.dir)", "불변식 3(파일 기반)"),
 ("kb/symbols.json", "에러 심볼 → 수정 라벨/대체/프로젝트/케이스 인덱스", "실행(harvest)", "트리아지 KB 매칭", "ledger와 함께 백업", ""),
 ("kb/cases/YYYY-MM/<case-id>.json", "KBCase: 성공한 수정 사례(시그니처, 패치, 라벨)", "실행(harvest)", "KB_PATCH 재생(LLM 없이), 레시피 후보 생성", "백업·공유", ""),
 ("kb/index.json", "L1/L2 시그니처 → 사례 인덱스(재생성 가능)", "실행", "KB 매칭", "지워져도 cases에서 재생성", ""),
 ("kb/recipes/candidates/ · active/ · rejected/", "학습된 레시피(yml + meta.json): 후보 → 검토 후 active → 이후 실행에 자동 포함", "실행(후보), 사람(승격/거절)", "kb.use_active_recipes일 때 rewrite에 포함", "active는 코드 리뷰처럼 검토 후 승격", ""),
 ("kb/patterns/*.yml", "손으로 쓴 패턴(시그니처 정규식 → 레시피/힌트/핸들러)", "개발자", "트리아지", "버전관리 가능", ""),
 ("fixtures/projects/<name>/", "소형 as-is 프로젝트 + README + expected.json(기대 결과)", "개발자", "scripts/e2e.sh, tests/fixtures", "새 시나리오마다 1개 추가", "24개"),
 ("fixtures/libs/", "사내 라이브러리 대역(ghost-lib, legacy-web-lib) + install.sh", "개발자", "e2e 전에 로컬 저장소에 설치", "", ""),
 ("fixtures/fleet/manifest.example.yml", "플릿 매니페스트 예시", "운영자가 복사", "sbmig fleet", "", ""),
 ("out/<run id>/profile.json", "분석 결과(ProjectProfile)", "analyze", "plan, 리포트, 사람", "실행 산출물(보관 선택)", ""),
 ("out/<run id>/plan.json", "마이그레이션 계획(스텝 목록·레시피·노트)", "plan", "실행 엔진, 사람", "", ""),
 ("out/<run id>/journal.jsonl", "이벤트 스트림(기계 판독)", "모든 단계", "웹 UI, 진단, 테스트", "문제 보고 시 첨부", ""),
 ("out/<run id>/sbmig.log", "DEBUG 로그(LLM 호출·빌드·결정; 키는 절대 없음)", "모든 단계", "진단", "문제 보고 시 첨부", ""),
 ("out/<run id>/logs/NNN-<label>.log", "빌드/리라이트 호출별 Maven 원본 로그(순번-라벨)", "build, rewrite", "로그 파서, 사람", "", "예: 006-build-verify.log"),
 ("out/<run id>/rewrite/", "합성된 rewrite.yml 등 OpenRewrite 입력", "rewrite", "Maven rewrite 플러그인", "", ""),
 ("out/<run id>/project/", "git 작업 사본(브랜치 sbmig/<run id>, 스텝·시도마다 커밋)", "Workspace", "모든 단계", "회사에서 git이 금지면 복사 기반 스냅샷이 후속 과제", ""),
 ("out/<run id>/changes.patch", "첫 커밋 대비 전체 diff(원본 바이트)", "finalize", "담당자 리뷰", "", ""),
 ("out/<run id>/result.zip", "마이그레이션된 프로젝트(.git 제외)", "finalize", "담당자에게 전달", "", ""),
 ("out/<run id>/report.md · report.json", "리포트(Handover 먼저)", "finalize", "담당자·운영자", "제출/보관", ""),
 ("out/<run id>/merge-conflicts/", "멀티모듈 merge에서 버려진 사본", "restructure", "수정 LLM 참고, 담당자", "", "라운드 7b/9"),
 ("out/<run id>/threads-*.txt", "워치독 자동 스레드 덤프(5분 무응답)", "fix loop watchdog", "개발자 진단", "", ""),
 ("out/web/<run id>/", "웹 UI 실행의 run 디렉터리(위와 동일 구조)", "웹 서버", "웹 UI", "web.max_finished_runs로 정리", ""),
 ("tests/", "오프라인 테스트(1,328개) + integration 마커 테스트", "개발자", "pytest", "", "docs/handbook 검증: tests/docs/test_handbook.py"),
 ("scripts/", "e2e.sh/e2e.py, fleet-demo.sh, demo-kb/run.sh", "개발자", "회귀 확인", "", ""),
 ("coding/", "같은 게이트웨이 계약 위의 독립 코딩 에이전트(sbmig 아님)", "개발자", "", "", "coding/README.md"),
 ("examples/adapters/", "LLM/수정기 플러그인 예시", "개발자", "llm.plugin / fix.plugin", "", "docs/ADAPTERS.md"),
 ("docs/", "영문 기술 문서(ARCHITECTURE, INTERFACES, ERROR-HANDLING, RUNBOOK, PLAN …)", "개발자", "", "", ""),
 ("docs/handbook/", "이 한국어 핸드북 + PPT + 엑셀", "개발자", "회사 담당자·운영자", "", ""),
]
ws, hr, n = sheet("디렉터리", "디렉터리별 저장 내용과 활용",
    "무엇이 저장되고, 누가 쓰고, 누가 읽고, 어떻게 다뤄야 하는지. 상세: docs/handbook/05-directories-and-storage.md",
    [("경로", "p"), ("저장 내용", "w"), ("쓰는 주체", "wr"), ("읽는 주체", "rd"), ("운영 메모(백업/정리)", "ops"), ("비고", "n")],
    [dict(zip(("p", "w", "wr", "rd", "ops", "n"), r)) for r in dirs], [40, 60, 22, 34, 44, 24])
specs.append(("디렉터리", "디렉터리별 저장·활용 요약", hr, n))

# ---- 목차 (first sheet) with COUNTA formulas
idx = wb.create_sheet("목차", 0)
idx["A1"] = "sbmig 레퍼런스 워크북"; idx["A1"].font = TITLE_FONT
idx["A2"] = ("JDK 21 + Spring Boot 4.1 마이그레이션 에이전트 sbmig의 표 중심 레퍼런스. 모든 표는 코드·설정·지식 파일에서 스크립트로 추출했습니다"
             "(생성 스크립트: 핸드북 빌드 시 사용; 수작업 표는 '디렉터리' 시트뿐). 본문 설명은 docs/handbook/*.md, 발표 요약은 sbmig-overview.pptx를 보십시오.")
idx["A2"].font = NOTE_FONT; idx.merge_cells("A2:D2"); idx["A2"].alignment = Alignment(wrap_text=True, vertical="top"); idx.row_dimensions[2].height = 45
for i, h in enumerate(["시트", "내용", "행 수(수식)", "짝이 되는 문서"], 1):
    c = idx.cell(row=4, column=i, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = BORDER
docmap = {"설정키": "08-config-reference.md", "핸들러": "04-triage-and-loop.md", "에러카탈로그": "09-operations.md", "CLI": "07-api-reference.md",
          "REST_API": "07-api-reference.md", "픽스처": "05-directories-and-storage.md", "저널이벤트": "01-pipeline.md", "리포트플래그": "09-operations.md",
          "레시피카탈로그": "02-workflows.md", "어드바이저리": "02-workflows.md", "지식노트": "04-triage-and-loop.md", "열거형": "README.md", "디렉터리": "05-directories-and-storage.md"}
for r, (name, desc, hr, n) in enumerate(specs, 5):
    idx.cell(row=r, column=1, value=name).font = BODY_FONT
    idx.cell(row=r, column=2, value=desc).font = BODY_FONT
    idx.cell(row=r, column=3, value=f"=COUNTA('{name}'!A{hr + 1}:A{hr + max(n, 1) + 5000})").font = BODY_FONT
    idx.cell(row=r, column=4, value=docmap.get(name, "")).font = BODY_FONT
    for c in range(1, 5):
        idx.cell(row=r, column=c).border = BORDER
last = 5 + len(specs)
idx.cell(row=last, column=2, value="합계").font = Font(name=FONT, bold=True, size=10)
idx.cell(row=last, column=3, value=f"=SUM(C5:C{last - 1})").font = Font(name=FONT, bold=True, size=10)
idx.cell(row=last + 2, column=1, value="범례: 모든 셀은 읽기 전용 참고용입니다. 회사 값(아티팩트 좌표, Nexus URL, LLM 키)은 이 파일에 없으며 config/artifacts.yml과 .env에만 둡니다.").font = NOTE_FONT
for col, w in zip("ABCD", (18, 60, 12, 34)):
    idx.column_dimensions[col].width = w
idx.freeze_panes = "A5"
wb.save(out)
print("saved", out, [ws.title for ws in wb.worksheets])
